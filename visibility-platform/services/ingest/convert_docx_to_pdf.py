"""One-off dev utility: render a .docx's text/table content into a plain
PDF using reportlab (already a project dependency, per pdf_report.py),
so it can flow through the existing PDF-only parsing.py unchanged.

Not part of the ingest pipeline itself -- a pre-processing step for
source documents that only exist as Word files. Preserves paragraph and
table-cell text verbatim (just reflowed), which is what verification.py's
verbatim-quote check needs downstream. Same pattern as check_db.py:
a standalone script, not a `python -m ingest.cli` subcommand, because
this is a manual one-time data-prep step, not an ingest capability.

Usage: python convert_docx_to_pdf.py <input.docx> <output.pdf>
"""
import sys

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table


def convert(input_path: str, output_path: str) -> None:
    d = docx.Document(input_path)
    styles = getSampleStyleSheet()
    story = []

    for block in _iter_block_items(d):
        if isinstance(block, docx.text.paragraph.Paragraph):
            text = block.text.strip()
            if text:
                # Escape reportlab's mini-markup special chars.
                safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, styles["Normal"]))
                story.append(Spacer(1, 4))
        else:  # Table
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            if rows:
                story.append(Table(rows))
                story.append(Spacer(1, 8))

    doc = SimpleDocTemplate(output_path, pagesize=A4)
    doc.build(story)


def _iter_block_items(document):
    """Yield paragraphs and tables in document order (python-docx doesn't
    expose this directly -- its .paragraphs/.tables are separate lists
    with no combined order)."""
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    parent_elm = document.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: python convert_docx_to_pdf.py <input.docx> <output.pdf>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
    print(f"wrote {sys.argv[2]}")
